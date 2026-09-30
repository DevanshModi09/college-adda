#!/usr/bin/env python3
"""
Convert the official JECRC CSE timetable workbook (one sheet per section) into the
JSON catalog the server imports.

    python3 server/scripts/xlsx_to_timetable.py "<workbook.xlsx>" server/data/catalog/<name>.json \
        --branch CSE --year 2 --term "III Sem 2026-27"

Sheet layout it expects (true for the 2026-27 III-sem workbook):
  - a header row whose first cell starts with "Day", followed by period columns
    like "Lecture 1 (8.00-8.50)";
  - one row per weekday (column A = "Monday"...), cells like "CN Lab-A VIB 502";
    merged cells span several periods (labs) or several days;
  - a legend below: "Computer Network(CN) | Mr X" (subject, faculty).
"""
import argparse
import json
import re
import sys

import openpyxl

DAYS = {"monday": 1, "tuesday": 2, "wednesday": 3, "thursday": 4, "friday": 5, "saturday": 6, "sunday": 0}
ROOM_RE = re.compile(r"\b(VIB|NYB)\s*-?\s*(\d{3}[A-Z]?)\b", re.I)
PERIOD_RE = re.compile(r"(\d{1,2})[.:](\d{2})\s*-\s*(\d{1,2})[.:](\d{2})")
# Section suffixes on subjects ("CN Lab-A", "COD-H", "OS-DG", "DSA Lab -H").
BATCH_RE = re.compile(r"\s*-\s*[A-Z]{1,2}$")
# Legend names without an "(ABBR)" and the abbreviation cells use for them.
ALIASES = {
    "software engineering": "SE",
    "prompt engineering": "Prompt Engg.",
    "r programming": "R Prog",
    "data structure(upgrade)": "Data Structure(upgrade)",
    "web development(upgrade)": "Web Development(upgrade)",
    "supervised learning": "Supervised Learning",
}


# Cell spellings that differ from the legend's.
SHORTHAND = {"prompt engg.": "PE", "prompt engg": "PE", "prompt engineering": "PE"}


def clean(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def to_24h(h: int, m: int) -> str:
    # The college day runs 8:00–17:00, so "1.30" means 13:30.
    if h < 8:
        h += 12
    return f"{h:02d}:{m:02d}"


def section_of(title: str):
    t = clean(title)
    m = re.match(r"^(?:Sec\s*)?([A-Z]{1,3})\s*(?:\((.*?)\)?)?\s*$", t, re.I)
    if not m:
        raise ValueError(f"can't read section from sheet title {title!r}")
    code, label = m.group(1).upper(), (m.group(2) or "").strip()
    label = re.sub(r"\s*\+\s*", " + ", label).strip(" +")
    return code, label


def merged_lookup(ws):
    """(row, col) -> (value, first_col, last_col) for every cell covered by a merge."""
    out = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = (v, rng.min_col, rng.max_col)
    return out


def parse_legend(ws, start_row):
    """abbr/lowercased name -> (full subject name, faculty)."""
    legend = {}
    for r in range(start_row, ws.max_row + 1):
        vals = [clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1)]
        vals = [v for v in vals if v]
        if not vals or vals[0].lower() in DAYS or vals[0].lower().startswith("subject"):
            continue
        name, faculty = vals[0], (vals[1] if len(vals) > 1 else "")
        m = re.match(r"^(.*?)\s*\(([^)]+)\)\s*$", name)
        if m and m.group(2).isupper():
            full, abbr = m.group(1).strip(), m.group(2).strip()
        else:
            full, abbr = name, ALIASES.get(name.lower(), name)
        if faculty.lower().startswith(("hod", "time-table")):
            faculty = ""
        legend[abbr.lower()] = (full, faculty)
        legend.setdefault(full.lower(), (full, faculty))
    return legend


def parse_cell(text: str, legend):
    text = clean(text)
    room_m = ROOM_RE.search(text)
    room = f"{room_m.group(1).upper()} {room_m.group(2)}" if room_m else ""
    subject = (text[: room_m.start()] + text[room_m.end():]) if room_m else text
    subject = subject.strip(" -")
    is_lab = bool(re.search(r"\blab\b", subject, re.I))
    base = re.sub(r"\blab\b", "", subject, flags=re.I)
    base = re.sub(r"\((?:Gen AI\+AWS\+MS|AIML)\)", "", base)
    base = BATCH_RE.sub("", base).strip(" -")
    base = re.sub(r"\s+SJ$", "", base)  # "Prompt Engg. SJ"
    base = re.sub(r"\s+\(", "(", base)  # "Data Structure (upgrade)" -> "Data Structure(upgrade)"
    base = SHORTHAND.get(base.lower(), base)

    full, faculty = legend.get(base.lower()) or next(
        ((f, fac) for k, (f, fac) in legend.items() if k.startswith(base.lower()) or base.lower().startswith(k)),
        (base, ""),
    )
    name = f"{full} Lab" if is_lab and not full.lower().endswith("lab") else full
    return {"subject": name, "code": base, "kind": "Lab" if is_lab else "Lecture", "room": room, "teacher": faculty}


def parse_sheet(ws):
    header = None
    for r in range(1, 10):
        if clean(ws.cell(r, 1).value).lower().startswith("day"):
            header = r
            break
    if header is None:
        raise ValueError(f"{ws.title}: no 'Day/Time' header row")

    periods = {}
    for c in range(2, ws.max_column + 1):
        m = PERIOD_RE.search(clean(ws.cell(header, c).value))
        if m:
            h1, m1, h2, m2 = map(int, m.groups())
            periods[c] = (to_24h(h1, m1), to_24h(h2, m2))
    if not periods:
        raise ValueError(f"{ws.title}: no period columns")

    day_rows, legend_start = {}, header + 1
    for r in range(header + 1, min(ws.max_row, header + 14) + 1):
        d = DAYS.get(clean(ws.cell(r, 1).value).lower())
        if d is not None:
            day_rows[r] = d
            legend_start = r + 1

    legend = parse_legend(ws, legend_start)
    merged = merged_lookup(ws)
    slots, seen = [], set()
    for r, day in day_rows.items():
        for c in sorted(periods):
            value, first, last = merged.get((r, c), (ws.cell(r, c).value, c, c))
            if not clean(value) or (r, first) in seen:
                continue
            seen.add((r, first))
            cols = [cc for cc in range(first, last + 1) if cc in periods]
            if not cols:
                continue
            slot = parse_cell(value, legend)
            slot.update(day=day, start=periods[cols[0]][0], end=periods[cols[-1]][1], raw=clean(value))
            slots.append(slot)
    return slots


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--branch", default="CSE")
    ap.add_argument("--year", type=int, default=2)
    ap.add_argument("--term", default="")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    sections = []
    for ws in wb.worksheets:
        code, label = section_of(ws.title)
        slots = parse_sheet(ws)
        sections.append({"key": f"{a.branch}|{a.year}|{code}", "code": code, "label": label, "slots": slots})
        print(f"{code:3} {label:32} {len(slots):3} slots", file=sys.stderr)

    with open(a.out, "w") as f:
        json.dump({"branch": a.branch, "year": a.year, "term": a.term, "source": a.xlsx.split("/")[-1].strip(), "sections": sections}, f, indent=1, ensure_ascii=False)
    print(f"wrote {len(sections)} sections to {a.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
